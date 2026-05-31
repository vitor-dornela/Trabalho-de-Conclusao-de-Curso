"""
Atualizar TCC v6 — preencher lacunas da Fundamentação + Metodologia:
  1. §56  — expandir explicação de Random Forest (ensemble, bagging, MDI)
  2. Nova seção 2.9  — MLBatLife + estratégia Greedy
  3. Novo § na Metodologia — justificativa regressão linear vs RF para EOL
"""
from pathlib import Path
from copy import deepcopy
from docx import Document
from docx.oxml.ns import qn
from lxml import etree
import shutil, sys

BASE = Path(__file__).resolve().parent.parent
SRC = BASE / "doc" / "TCC_Modelo_Estrutural_Formatado_v6.docx"
BAK = BASE / "doc" / "_old" / "TCC_v6_pre_gaps.docx"

# ── Backup ───────────────────────────────────────
BAK.parent.mkdir(parents=True, exist_ok=True)
shutil.copy2(SRC, BAK)
print(f"Backup → {BAK.name}")

doc = Document(str(SRC))

# ── Helpers ──────────────────────────────────────

def find_para(prefix):
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith(prefix):
            return i, p
    raise ValueError(f"Não encontrado: '{prefix[:60]}...'")


def replace_text(para, new_text):
    runs = para.runs
    if runs:
        runs[0].text = new_text
        for r in runs[1:]:
            r._element.getparent().remove(r._element)
    else:
        para.text = new_text


def clone_para(template, text):
    """Clone paragraph XML from template, replace text."""
    new_p = deepcopy(template._element)
    # Remove runs
    for r in new_p.findall(qn("w:r")):
        new_p.remove(r)
    # Remove bookmarks (avoid duplicate IDs)
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for bm in new_p.findall(qn(tag)):
            new_p.remove(bm)
    # New run
    r_el = etree.SubElement(new_p, qn("w:r"))
    src_runs = template._element.findall(qn("w:r"))
    if src_runs:
        rPr = src_runs[0].find(qn("w:rPr"))
        if rPr is not None:
            r_el.insert(0, deepcopy(rPr))
    t_el = etree.SubElement(r_el, qn("w:t"))
    t_el.text = text
    t_el.set(qn("xml:space"), "preserve")
    return new_p


# ═════════════════════════════════════════════════
# 1) Substituir §56 — Random Forest expandido
# ═════════════════════════════════════════════════
i56, p56 = find_para("Modelos como Random Forest, Gradient Boosting")

replace_text(
    p56,
    "Diversos algoritmos de aprendizado de máquina têm sido empregados "
    "na estimativa de SOH e EOL, incluindo Random Forest, Gradient "
    "Boosting, Support Vector Regression, Gaussian Process Regression "
    "e redes neurais (RASMUSSEN; WILLIAMS, 2006; LUQUE et al., 2025). "
    "Entre esses, o Random Forest (BREIMAN, 2001) destaca-se por sua "
    "aplicação em problemas de regressão com dados tabulares. O "
    "algoritmo constrói um conjunto (ensemble) de árvores de decisão, "
    "cada uma treinada sobre uma amostra aleatória com reposição "
    "(bootstrap) dos dados originais. A predição final é obtida pela "
    "média das saídas de todas as árvores, o que reduz a variância e "
    "aumenta a robustez do modelo. Os principais hiperparâmetros são o "
    "número de estimadores (n_estimators) e a profundidade máxima de "
    "cada árvore (max_depth). Além disso, o Random Forest permite "
    "quantificar a importância de cada variável de entrada por meio da "
    "diminuição média de impureza (Mean Decrease in Impurity \u2014 MDI), "
    "que mede a contribuição de cada feature para a redução do erro nas "
    "divisões internas das árvores (BREIMAN, 2001; PEDREGOSA et al., 2011).",
)
print(f"[1/3] §{i56} \u2014 RF expandido \u2713")

# ═════════════════════════════════════════════════
# 2) Inserir seção 2.9 após §56
# ═════════════════════════════════════════════════
# Heading 2 template
h2 = None
for p in doc.paragraphs:
    if p.style and p.style.name == "Heading 2":
        h2 = p
        break
if not h2:
    raise ValueError("Heading 2 template não encontrado")

# Inserir body DEPOIS heading (addnext insere logo após, ordem reversa)
body_el = clone_para(
    p56,
    "O MLBatLife (LUQUE et al., 2025) é um modelo preditivo pré-treinado "
    "baseado em Random Forest, desenvolvido para estimar a degradação de "
    "baterias em sistemas fotovoltaicos residenciais a partir de perfis "
    "de potência residual. O modelo integra a extração de variáveis "
    "diárias e a estimativa da perda de capacidade em um pipeline "
    "automatizado, eliminando a necessidade de engenharia de features "
    "manual. Entre as estratégias de operação da bateria disponíveis no "
    "modelo, a estratégia Greedy prioriza o autoconsumo imediato da "
    "energia excedente, carregando e descarregando a bateria sem limitar "
    "a profundidade de descarga. Essa estratégia representa o "
    "comportamento padrão de sistemas residenciais sem gerenciamento "
    "otimizado e tende a impor ciclos mais intensos à bateria em "
    "comparação com estratégias que restringem a profundidade de "
    "descarga ou otimizam o despacho (LUQUE et al., 2025).",
)
p56._element.addnext(body_el)

head_el = clone_para(h2, "O modelo MLBatLife e estratégias de operação")
p56._element.addnext(head_el)
# Resultado: §56 → heading_2.9 → body_2.9

print("[2/3] Seção 2.9 inserida \u2014 MLBatLife + Greedy \u2713")

# ═════════════════════════════════════════════════
# 3) Justificativa regressão linear na METODOLOGIA
# ═════════════════════════════════════════════════
i64, p64 = find_para("Na etapa de modelagem, o MLBatLife")

reg_el = clone_para(
    p64,
    "Para a estimativa do fim de vida útil (EOL), optou-se pela "
    "regressão linear sobre a curva de SOH em vez de estender as "
    "predições do Random Forest além do período observado. Essa escolha "
    "decorre de uma limitação inerente a modelos baseados em árvores de "
    "decisão: por operarem mediante partições do espaço de features, "
    "esses modelos não extrapolam tendências para valores de entrada "
    "fora do domínio de treinamento, retornando a predição do nó folha "
    "mais próximo (BREIMAN, 2001). Como a degradação da bateria após o "
    "período inicial de formação da camada SEI segue um regime "
    "aproximadamente linear (BARRÉ et al., 2013), a regressão linear "
    "permite projetar o EOL de forma confiável, ancorando a projeção "
    "no último SOH observado.",
)
p64._element.addnext(reg_el)

print(f"[3/3] Justificativa regressão após §{i64} \u2713")

# ═════════════════════════════════════════════════
# Salvar
# ═════════════════════════════════════════════════
try:
    doc.save(str(SRC))
    print(f"\nSalvo: {SRC.name}")
except PermissionError:
    print("\n\u26a0 ERRO: feche o Word antes de salvar!", file=sys.stderr)
    sys.exit(1)

# ═════════════════════════════════════════════════
# Verificação
# ═════════════════════════════════════════════════
doc2 = Document(str(SRC))
print("\n\u2500\u2500 Fundamentação (§51\u2192§62) \u2500\u2500\n")
for i, p in enumerate(doc2.paragraphs):
    t = p.text.strip()
    s = p.style.name if p.style else "?"
    if 51 <= i <= 62 and t:
        marker = " <<<" if s == "Heading 2" else ""
        print(f"  {i:3d} [{s}] {t[:120]}{'…' if len(t) > 120 else ''}{marker}")

print("\n\u2500\u2500 Metodologia (§64\u2192§72) \u2500\u2500\n")
for i, p in enumerate(doc2.paragraphs):
    t = p.text.strip()
    s = p.style.name if p.style else "?"
    if 64 <= i <= 72 and t:
        print(f"  {i:3d} [{s}] {t[:120]}{'…' if len(t) > 120 else ''}")
